from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from config import MAPA_ETAPAS  # noqa: E402
from main import gerar_relatorio  # noqa: E402
from src.tarefas import extrair_tarefas  # noqa: E402


RESPOSTAS_HEADERS = [
    "Id",
    "Hora de início",
    "Hora de conclusão",
    "Email",
    "Nome",
    "obra",
    "nome2",
    "data",
    "tempo_manha",
    "interrupcao_manha",
    "tempo_tarde",
    "interrupcao_tarde",
    "mestre_de_Obras",
    "eletricista",
    "pedreiro",
    "servente",
    "encanador",
    "pintor",
    "etapa",
    "fundação",
    "estrutura",
    "alvenaria",
    "cobertura",
    "instalações",
    "acabamento",
    "outros_procedimentos",
    "ocorrencias",
]

CADASTRO_HEADERS = [
    "obra_id",
    "objeto",
    "contrato",
    "contratante",
    "endereco",
    "fiscal",
    "Crea-MT",
]


def criar_base_teste(caminho: Path) -> None:
    workbook = Workbook()

    respostas = workbook.active
    respostas.title = "resposta_forms"
    respostas.append(RESPOSTAS_HEADERS)
    respostas.append(
        [
            1,
            datetime(2026, 5, 21, 7, 30),
            datetime(2026, 5, 21, 17, 0),
            "teste@example.com",
            "Usuario Teste",
            "Obra_1 - Projeto teste",
            "Equipe A",
            datetime(2026, 5, 21),
            "Limpo",
            "",
            "Nublado",
            "",
            "1",
            "0",
            "2",
            "3",
            "0",
            "0",
            "Outros",
            "",
            "",
            "",
            "",
            "",
            "",
            "Limpeza de canteiro;Organizacao de estoque;",
            "Registro criado pelo smoke test.",
        ]
    )

    cadastro = workbook.create_sheet("cadastro_obras")
    cadastro.append(CADASTRO_HEADERS)
    cadastro.append(
        [
            "Obra_1",
            "Projeto teste",
            "CT-001",
            "Contratante Teste",
            "Rua Teste",
            "Fiscal Teste",
            "000000",
        ]
    )

    workbook.save(caminho)


def verificar_tarefas_outros() -> None:
    registro = dict(
        zip(
            RESPOSTAS_HEADERS,
            [
                1,
                None,
                None,
                "",
                "",
                "Obra_1 - Projeto teste",
                "",
                datetime(2026, 5, 21),
                "",
                "",
                "",
                "",
                0,
                0,
                0,
                0,
                0,
                0,
                "Outros",
                "",
                "",
                "",
                "",
                "",
                "",
                "Limpeza de canteiro;Organizacao de estoque;",
                "",
            ],
        )
    )

    tarefas = extrair_tarefas(registro, MAPA_ETAPAS)
    if tarefas != ["Limpeza de canteiro", "Organizacao de estoque"]:
        raise SystemExit(f"Falha ao extrair tarefas da etapa Outros: {tarefas}")


def verificar_docx(caminho: Path) -> None:
    if not caminho.exists() or caminho.stat().st_size == 0:
        raise SystemExit("DOCX de smoke test nao foi gerado.")

    placeholders = []
    with ZipFile(caminho) as docx:
        for nome in docx.namelist():
            if not nome.startswith("word/") or not nome.endswith(".xml"):
                continue
            texto = docx.read(nome).decode("utf-8", errors="ignore")
            if "{{" in texto or "{%" in texto:
                placeholders.append(nome)

    if placeholders:
        raise SystemExit(f"Placeholders nao renderizados no DOCX: {placeholders}")


def main() -> None:
    verificar_tarefas_outros()

    pasta_saida = ROOT / "build" / "ci-smoke"
    pasta_saida.mkdir(parents=True, exist_ok=True)

    caminho_excel = pasta_saida / "base_teste.xlsx"
    caminho_docx = pasta_saida / "diario_teste.docx"

    criar_base_teste(caminho_excel)

    gerar_relatorio(
        obra="Obra_1 - Projeto teste",
        data_inicio_txt="21/05/2026",
        data_fim_txt="21/05/2026",
        medicao="1",
        data_assinatura="21/05/2026",
        caminho_saida=caminho_docx,
        assumir_ultimo_duplicado=True,
        caminho_excel=caminho_excel,
        caminho_template=ROOT / "templates" / "modelopadrao.docx",
    )

    verificar_docx(caminho_docx)
    print(f"Smoke test gerou {caminho_docx.relative_to(ROOT)} com sucesso.")


if __name__ == "__main__":
    main()
